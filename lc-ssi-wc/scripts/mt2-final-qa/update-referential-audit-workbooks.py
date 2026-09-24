from __future__ import annotations

import hashlib
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
FILES = (
    ROOT / "qa/tdd/mt2/MT2XX_測試案例_SSI與NOSTRO_v6.xlsx",
    ROOT / "qa/fixtures/mt2/MT2XX_測試案例_SSI與NOSTRO_v6.xlsx",
)
EXPECTED_SOURCE_SHA256 = {
    "49A0B3B82D5C0CE5B6FE791B30E6122E7B1838070AB7E153B2BEE49C2DB50EA7",
    "786F9CF122FBD15ED5F0CB1DD1C6ECA7376B08B1794AB3E90C8B59FF4E85EC8F",
    "5FF5DB39AA0D54C85AC4E002139A0A5C3D194E5B229EFD6B0FEC19A7A2E4149C",
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def row_for(sheet, first_column_value: str, column: int = 1) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, column).value == first_column_value:
            return row
    raise ValueError(f"Missing row {first_column_value!r} in {sheet.title}")


def update(source: Path, output: Path) -> None:
    workbook = load_workbook(source)

    audit = workbook["Source Fidelity Audit"]
    audit_row = row_for(audit, "ACTIVE SSI accountId without ACTIVE Nostro join")
    audit.cell(audit_row, 2).value = 0
    audit.cell(audit_row, 3).value = 42
    audit.cell(audit_row, 4).value = "[]"
    audit.cell(audit_row, 5).value = (
        "Exact equality SSI.route.accountId = ACTIVE Nostro.accountReference; "
        "id and maskedAccountRef are not join keys. Executable cross-check also "
        "matches accountServicerBic, currency, SETTLEMENT purpose, booking entity "
        "and effective date."
    )

    gaps = workbook["Data Gaps"]
    gap_row = row_for(gaps, "SSI → Own Nostro referential integrity", column=2)
    gaps.cell(gap_row, 1).value = "RESOLVED DATA FACT / HARDENING OPEN"
    gaps.cell(gap_row, 3).value = (
        "0/42 orphan. All 42 ACTIVE SSI route.accountId values exactly match at "
        "least one ACTIVE Nostro.accountReference; all 42 also have an eligible "
        "servicer/currency/entity/effective-date composite as of 2026-09-10."
    )
    gaps.cell(gap_row, 4).value = (
        "Preserve exact accountReference equality and PROFILE_INCOMPLETE "
        "fail-closed behavior. An immutable nostroId foreign key remains optional "
        "schema hardening; never substitute id, maskedAccountRef, prefix or fuzzy matching."
    )
    account_reference_area = "Nostro accountReference completeness"
    try:
        reference_row = row_for(gaps, account_reference_area, column=2)
    except ValueError:
        reference_row = gaps.max_row + 1
        template_row = reference_row - 1
        for column in range(1, gaps.max_column + 1):
            source_cell = gaps.cell(template_row, column)
            target_cell = gaps.cell(reference_row, column)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)
    gaps.cell(reference_row, 1).value = "DATA GOVERNANCE"
    gaps.cell(reference_row, 2).value = account_reference_area
    gaps.cell(reference_row, 3).value = (
        "69/159 ACTIVE Nostro rows omit accountReference. Independent snapshot "
        "verification confirms that 68 of them also omit allowedBookingEntities, "
        "but this does not orphan "
        "any of the current 42 ACTIVE SSI rows."
    )
    gaps.cell(reference_row, 4).value = (
        "Backfill accountReference under controlled data governance. Keep "
        "maskedAccountRef display-only; never use prefix, substring or fuzzy matching."
    )

    cases = workbook["Test Cases"]
    case_row = row_for(cases, "MT202-34")
    request = json.loads(cases.cell(case_row, 3).value)
    request["MRG business scenario"] = (
        "Controlled negative fixture whose SSI accountId has no exact ACTIVE "
        "Nostro.accountReference"
    )
    request["MRG evidence"] = (
        "Negative referential-integrity contract; authoritative snapshot has 0 "
        "natural orphan; controlled fixture QA-SSI-ORPHAN-001 is required"
    )
    request["dataPrecondition"] = "CONTROLLED_ISOLATED_FIXTURE"
    request["request/context"]["candidate"] = {
        "ssiCode": "QA-SSI-ORPHAN-001",
        "accountId": "DEMO-NOSTRO-NONEXISTENT",
        "nostroMatch": None,
    }
    request["current DB fixtures"] = [
        "Synthetic isolated orphan required; no matching ACTIVE Nostro.accountReference"
    ]
    cases.cell(case_row, 3).value = json.dumps(request, ensure_ascii=False, indent=2)
    mt_expected = json.loads(cases.cell(case_row, 5).value)
    mt_expected["missingRelationship"] = (
        "SSI.route.accountId -> ACTIVE Nostro.accountReference"
    )
    cases.cell(case_row, 5).value = json.dumps(
        mt_expected, ensure_ascii=False, indent=2
    )

    workbook.save(output)


def main() -> None:
    hashes = {path: digest(path) for path in FILES}
    if len(set(hashes.values())) != 1:
        raise RuntimeError(f"Workbook copies diverged before update: {hashes}")
    before_hash = next(iter(hashes.values()))
    if before_hash not in EXPECTED_SOURCE_SHA256:
        raise RuntimeError(f"Unexpected source workbook SHA-256: {hashes}")
    temporary = FILES[0].with_suffix(".referential-audit.tmp.xlsx")
    update(FILES[0], temporary)
    payload = temporary.read_bytes()
    for destination in FILES:
        destination.write_bytes(payload)
    temporary.unlink()
    print(
        json.dumps(
            {
                "status": "UPDATED",
                "beforeSha256": before_hash,
                "afterSha256": digest(FILES[0]),
                "copiesIdentical": digest(FILES[0]) == digest(FILES[1]),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
