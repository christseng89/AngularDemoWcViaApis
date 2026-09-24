from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import sqlite3
from pathlib import Path

import openpyxl

UI_FAIL_CLOSED_SHEET = "UI Fail-Closed"
SCENARIO_BEHAVIOR_SHEET = "Scenario Behavior"
MESSAGES = ("MT202", "MT205", "MT202COV", "MT205COV")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def route_value(payload: dict, key: str):
    value = payload.get(key)
    return value if value is not None else (payload.get("route") or {}).get(key)


def copy_row_style(sheet, source_row: int, target_row: int, columns: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, columns + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)


def load_active_ssi(snapshot: Path) -> dict[str, dict]:
    uri = f"file:{snapshot.as_posix()}?mode=ro&immutable=1"
    connection = sqlite3.connect(uri, uri=True)
    try:
        records = [json.loads(row[0]) for row in connection.execute("select payload from ssi")]
    finally:
        connection.close()
    active = {}
    for record in records:
        code = route_value(record, "ssiCode")
        if record.get("status") == "ACTIVE" and code:
            active[code] = record
    return active


def case_note(pair: dict, chosen: str | None) -> str:
    if chosen is None:
        return (
            "受控 top-rank tie：SSI-DEMO-003／022 同為 priority 10、PRIMARY、"
            "NAMED_COUNTERPARTY。必須 fail closed；chosenRoute／agent roles／"
            "roleProvenance／messageComposerContext 為 null，candidates 無序，payloadGenerated=false。"
        )
    if len(pair["ssiCodes"]) == 1:
        return "唯一候選；須驗 chosenRoute、roleProvenance、snapshotHash、resolutionToken 與幣別一致。"
    return (
        f"{len(pair['ssiCodes'])} 筆 eligible；依受控 rank keys 選 {chosen}。"
        "須列出 alternatives，且帳號、agent、provenance 均來自同一 chosen SSI。"
    )


def plan_bank_names(plan) -> dict[str, str]:
    return {
        plan.cell(row, 4).value: plan.cell(row, 5).value
        for row in range(5, 61)
        if plan.cell(row, 4).value
    }


def configure_execution_columns(plan) -> None:
    execution_headers = [
        "Resolver Baseline Profile", "Counterparty Bank Service ID",
        "Actual HTTP", "Actual code", "Actual chosen SSI / version",
        "Actual candidates", "Snapshot Hash", "Resolution Token", "Correlation / Request ID",
        "Tester", "Execution Date", "PASS / FAIL", "Evidence Path", "Screenshot / Defect ID",
    ]
    for column, header in enumerate(execution_headers, 17):
        plan.cell(4, column).value = header
        copy_row_style(plan, 4, 4, column)
        plan.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 23
    for row in range(5, 53):
        copy_row_style(plan, 6, row, 30)


def plan_case_values(
    pair: dict,
    pair_number: int,
    message: str,
    case_number: int,
    bank_names: dict[str, str],
    allow_list: dict,
) -> list:
    ambiguous = pair["counterpartyBic"] == "BARCGB22" and pair["currency"] == "GBP"
    cov = message.endswith("COV")
    chosen = None if ambiguous else pair["ssiCodes"][0]
    return [
        f"UAT-V151-{case_number:03d}", pair_number, "EXECUTABLE NOW",
        pair["counterpartyBic"], bank_names.get(pair["counterpartyBic"], pair["counterpartyBic"]),
        pair["currency"], pair["bookingEntity"], message,
        allow_list["messageDefinitionId"],
        allow_list["businessService"] if cov else "swift.cbprplus.04",
        422 if ambiguous else 200,
        "SSI_AMBIGUOUS" if ambiguous else "SSI_RESOLVED",
        "null（fail closed）" if ambiguous else chosen,
        len(pair["ssiCodes"]), "留白（由 eligible route 推導）", case_note(pair, chosen),
        f"RESOLVER_BASELINE_{message}", f"BANK-SVC-{pair['counterpartyBic']}",
        None, None, None, None, None, None, None,
        None, None, None, None, None,
    ]


def populate_plan_cases(plan, allow_list: dict, bank_names: dict[str, str]) -> None:
    case_number = 1
    target_row = 5
    for pair_number, pair in enumerate(allow_list["pairs"], 1):
        for message in MESSAGES:
            values = plan_case_values(
                pair, pair_number, message, case_number, bank_names, allow_list
            )
            for column, value in enumerate(values, 1):
                plan.cell(target_row, column).value = value
            target_row += 1
            case_number += 1


def clear_unused_plan_rows(plan) -> None:
    for row in range(53, 61):
        for column in range(1, 31):
            plan.cell(row, column).value = None


def configure_plan_summary(plan, snapshot: Path) -> None:
    plan["A1"] = "MT2xx SSI Resolution — v15.1 現行資料 UAT 執行清單"
    plan["A2"] = (
        f"資料來源：post-migration WAL-aware snapshot {snapshot.name}（SHA-256 {sha256(snapshot)}）｜"
        "12 組 × 4 電文 = 48 案｜預期 44 RESOLVED + 4 SSI_AMBIGUOUS"
    )
    plan["A62"] = "彙總（依受控矩陣產生；非公式，任何閱讀器皆可見）"
    for row, value in {63: 48, 64: 48, 65: 0, 66: 44, 67: 4, 68: 0, 69: 0}.items():
        plan.cell(row, 2).value = value
    plan["A65"] = "NOT_EXECUTABLE"
    plan["A68"] = "預期 503 PROFILE_INCOMPLETE（矩陣內）"
    plan["A69"] = "NOT_EXECUTABLE 比率"
    plan["A71"] = (
        "矩陣外負向控制：NSSIUSN1／USD／MT202COV 與 MT205COV 必須在 rank 前回 HTTP 503 PROFILE_INCOMPLETE，"
        "payloadGenerated=false，且不得 fallback 至 Core。此控制不計入 48 案功能矩陣。"
    )
    plan.auto_filter.ref = "A4:AD52"
    plan.freeze_panes = "D5"


def populate_plan(plan, allow_list: dict, snapshot: Path) -> None:
    bank_names = plan_bank_names(plan)
    configure_execution_columns(plan)
    populate_plan_cases(plan, allow_list, bank_names)
    clear_unused_plan_rows(plan)
    configure_plan_summary(plan, snapshot)


def refresh_candidates(candidates, allow_list: dict, active: dict[str, dict]) -> None:
    candidates.cell(4, 24).value = "businessService"
    candidates.cell(4, 25).value = "sourceMessageTypes"
    for column in (24, 25):
        copy_row_style(candidates, 4, 4, column)
        candidates.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 28
    allowed_codes = [code for pair in allow_list["pairs"] for code in pair["ssiCodes"]]
    row_by_code = {
        candidates.cell(row, 1).value: row
        for row in range(5, candidates.max_row + 1)
        if candidates.cell(row, 1).value
    }
    for code, payload in active.items():
        row = row_by_code.get(code)
        if row:
            candidates.cell(row, 2).value = payload.get("id")
            candidates.cell(row, 3).value = payload.get("version")
    for code in allowed_codes:
        row = row_by_code[code]
        payload = active[code]
        candidates.cell(row, 24).value = route_value(payload, "businessService")
        candidates.cell(row, 25).value = route_value(payload, "sourceMessageTypes")
    for code in ("SSI-DEMO-041", "SSI-DEMO-042"):
        row = row_by_code.get(code)
        if row:
            candidates.cell(row, 9).value = "NEGATIVE CONTROL — NOT COV ALLOW-LISTED"
            candidates.cell(row, 24).value = route_value(active[code], "businessService")
            candidates.cell(row, 25).value = route_value(active[code], "sourceMessageTypes")


def refresh_guidance(workbook) -> None:
    guide = workbook["輸入指引"]
    guide["B5"] = "依案例：MT 或 MX；切換格式不得重新解析或改變 confirmed snapshot"
    guide["C7"] = (
        "硬過濾。任一 resolver input 變動後，舊結果須立即失效；重新 Preview 後 currency、chosen SSI、"
        "agent、nostro 與 provenance 必須整批一致。"
    )


def refresh_evidence(
    workbook,
    root: Path,
    snapshot: Path,
    report_path: Path,
    allow_list_path: Path,
    report: dict,
) -> None:
    evidence = workbook["證據與範圍"]
    evidence["B5"] = str(snapshot.relative_to(root)).replace("\\", "/")
    evidence["B6"] = sha256(snapshot)
    evidence["B7"] = snapshot.stat().st_size
    evidence["B12"] = "post-migration 全量 scan：12 組中僅 BARCGB22／GBP／HK01 top-rank tie；四種 MT2xx 電文均預期 422。"
    evidence["B13"] = "12 組、MT202／MT205／MT202COV／MT205COV，共 48 案；全部使用 pacs.009.001.08 canonical tuple。"
    evidence["A15"] = "v15.1 migration"
    evidence["B15"] = "24 筆 SSI 具 explicit COV allow-list；Core/COV 以 businessService 分流，禁止 synthetic .COV messageDefinitionId。"
    evidence["A16"] = "現行驗收帳本"
    evidence["B16"] = "48 = 44 SSI_RESOLVED + 4 SSI_AMBIGUOUS；NOT_EXECUTABLE=0；另有 1 筆 Northstar COV 503 負向控制。"
    evidence["A17"] = "最新自動執行證據"
    evidence["B17"] = (
        f"{report_path.relative_to(root).as_posix()}；SHA-256 {sha256(report_path)}；"
        f"passed={report['totals']['passed']}/{report['totals']['planned']}。"
    )
    evidence["A18"] = "BA review gate"
    evidence["B18"] = "候選版須經獨立 BA 覆核接受後才可置入 qa/reports/latest/mt2；不得以自動執行成功取代業務驗收。"
    evidence["A19"] = "受控 v15.1 baseline"
    evidence["B19"] = "qa/tdd/mt2/MT2XX_支援標準SSI_SR2026_MRG與ISO20022對應覆核版_v15.1.xlsx；SHA-256 0493C1D44D55A48B369AF6C6198BEE6943C1E17FC3A76AC93DE22C1D2A69A3D5。"
    evidence["A20"] = "COV allow-list"
    evidence["B20"] = f"parameters/cov-profile-allow-list.v15.1.json；SHA-256 {sha256(allow_list_path)}。"


def refresh_scenario_reference(workbook) -> None:
    scenario_reference = workbook["情境對照"]
    scenario_reference["F8"] = (
        "12 組 approved allow-list 已可執行：11 組 RESOLVED；BARCGB22/GBP 回 SSI_AMBIGUOUS。"
        "僅未列入 allow-list 的 Northstar 負向控制回 503 PROFILE_INCOMPLETE。"
    )


def replace_sheet(workbook, title: str):
    if title in workbook.sheetnames:
        del workbook[title]
    return workbook.create_sheet(title)


def style_table(sheet, plan, widths: tuple[int, ...]) -> None:
    for cell in sheet[1]:
        cell.font = copy.copy(plan["A4"].font)
        cell.fill = copy.copy(plan["A4"].fill)
        cell.alignment = copy.copy(plan["A4"].alignment)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)


def add_ui_fail_closed_sheet(workbook, plan) -> None:
    ui = replace_sheet(workbook, UI_FAIL_CLOSED_SHEET)
    ui.append(["Case ID", "Precondition", "Action", "Expected Result", "Business Risk", "Execution Record"])
    ui_cases = [
        (
            "UI-STALE-01", "BARCGB22 / USD 已 SSI_RESOLVED",
            "改 Currency 為 SGD（No SSI Record），不按 Preview",
            "舊結果立即清空或明確標 stale；不得顯示 RESOLVED、舊 agent、nostro 或 token。",
            "避免使用者把 USD 路由誤當 SGD 路由", "",
        ),
        (
            "UI-RACE-01", "先送 USD Preview，立即切 GBP 再送 Preview",
            "以瀏覽器 route interception 固定延遲 USD response 1500 ms；GBP response 先返回",
            "畫面只接受最新 GBP request identity；遲到 USD response 不得回填。",
            "避免非同步競態造成跨幣別路由混用", "",
        ),
        (
            "UI-BANK-SVC-01", "未選 Counterparty Bank Service",
            "以空 bankServiceId 送出 Preview",
            "fail closed；不得以 raw BIC fallback，須顯示可行動的欄位錯誤。",
            "維持 bankServiceId → BIC 的唯一受控解析鏈", "",
        ),
        (
            "UI-BANK-SVC-02", "輸入不存在的 Bank Service ID",
            "送出 Preview",
            "HTTP 422 OPTION_CONSTRAINT_VIOLATION；不得產生 payload，亦不得接受 raw BIC fallback。",
            "防止未登錄銀行身分繞過 Bank Service", "",
        ),
    ]
    for case in ui_cases:
        ui.append(case)
    ui.freeze_panes = "A2"
    ui.auto_filter.ref = f"A1:F{ui.max_row}"
    style_table(ui, plan, (18, 32, 40, 56, 38, 28))


def add_scenario_behavior_sheet(workbook, plan) -> None:
    scenarios = replace_sheet(workbook, SCENARIO_BEHAVIOR_SHEET)
    scenarios.append(["Case ID", "Message", "Target Scenario", "Expected Behavior", "Separate From 48-Case Resolver Matrix"])
    scenario_rows = [
        ("SCN-202-01", "MT202", "BOOK_TRANSFER_SAME_RECEIVER", "驗證同 receiver 的 57a 省略語意與 MT/MX 呈現差異。", "YES"),
        ("SCN-202-02", "MT202", "CREDIT_ONE_OF_SEVERAL_AT_57A", "驗證明示帳戶、57a 與 chosen SSI accountId 一致。", "YES"),
        ("SCN-205-01", "MT205", "INITIAL_MT200_201_EQUIVALENCE", "驗證前手 context 與 mandatory 52a；缺值回 MESSAGE_CONTEXT_MISSING。", "YES"),
        ("SCN-205COV-01", "MT205COV", "NO_MT200_201_EQUIVALENCE", "驗證 cover context 傳遞且不套用 MT205-only 規則。", "YES"),
        ("SCN-202COV-GAP", "MT202COV", "NO_REGISTERED_UI_SCENARIO", "登錄為情境覆蓋缺口；只執行 resolver baseline，不宣稱 scenario coverage。", "YES"),
    ]
    for row in scenario_rows:
        scenarios.append(row)
    scenarios["A8"] = "說明"
    scenarios["B8"] = "48 案主矩陣的 Resolver Baseline Profile 不是可選 UI scenario；不需套用情境。以上 5 案另計，不灌入 48 案分母。"
    scenarios.freeze_panes = "A2"
    scenarios.auto_filter.ref = f"A1:E{scenarios.max_row}"
    style_table(scenarios, plan, (20, 14, 40, 68, 24))


def prepare(source: Path, output: Path, root: Path) -> None:
    allow_list_path = root / "parameters" / "cov-profile-allow-list.v15.1.json"
    snapshot = root / "qa" / "mt2-final" / "fixtures" / "baselines" / "ssi-demo.v15.1-post-migration.sqlite"
    report_path = root / "qa" / "reports" / "MT2XX_v15.1_post_migration_tie_scan_20260911.json"
    allow_list = json.loads(allow_list_path.read_text(encoding="utf-8"))
    report = json.loads(report_path.read_text(encoding="utf-8"))
    active = load_active_ssi(snapshot)
    workbook = openpyxl.load_workbook(source)
    plan = workbook["UAT 執行清單"]

    populate_plan(plan, allow_list, snapshot)
    refresh_candidates(workbook["SSI 候選明細"], allow_list, active)
    refresh_guidance(workbook)
    refresh_evidence(workbook, root, snapshot, report_path, allow_list_path, report)
    refresh_scenario_reference(workbook)
    add_ui_fail_closed_sheet(workbook, plan)
    add_scenario_behavior_sheet(workbook, plan)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    prepare(args.source.resolve(), args.output.resolve(), args.root.resolve())
    print(json.dumps({"output": str(args.output.resolve()), "sha256": sha256(args.output.resolve()), "bytes": args.output.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
