"""Build the governed MT347 catalogue and QA/UAT fixture artifacts from TDD v5.

The workbook is the approved contract.  Positive fixtures are suitable for the
canonical Development/Demo seed.  Negative fixtures remain isolated and must
only be loaded by the negative test runner.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sqlite3
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCOPE = {
    "MT300", "MT304", "MT305", "MT306", "MT320", "MT330", "MT340",
    "MT341", "MT350", "MT360", "MT361", "MT362", "MT364", "MT365",
    "MT400", "MT730", "MT734", "MT742", "MT750", "MT752", "MT754",
    "MT756", "MT765", "MT768", "MT769",
}
WORKBOOK_SHA256 = "82C6ABCFD91E7D35E1382F8C86BF796D8DBF05CF8C8D94F9F0C7B5889AB52C9A"
MEMORY_SHA256 = "682075A1EC36AD1A5397382063B1AFEE08F0E88921E139E7E62A8033552A985A"
CASE_ID_COLUMN = "Test Case No."
CATALOGUE_FILE = "parameters/ssi-mappings.sr2026.json"
MT347_MAKER = "maker.mt347"
MT347_CHECKER = "checker.mt347"


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def rows_by_header(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    header = [str(value) if value is not None else "" for value in rows[0]]
    return [dict(zip(header, row, strict=True)) for row in rows[1:] if row[0]]


def parse_options(value: Any) -> set[str]:
    return {part.strip().upper() for part in str(value).split("/") if part.strip()}


def parse_nvr_refs(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text or text.lower() == "none":
        return []
    return sorted(set(re.findall(r"\b(?:C\d+|D\d+|E\d+|T\d+)\b", text)))


def governed_rules(workbook: Any) -> dict[tuple[str, str, str], dict[str, Any]]:
    result: dict[tuple[str, str, str], dict[str, Any]] = {}
    for row in rows_by_header(workbook, "Confirmed Rules"):
        message = str(row["Message Type"])
        if message not in SCOPE:
            continue
        tag = str(row["Tag"])[:2]
        key = (message, str(row["Sequence"]), tag)
        result[key] = {
            "options": parse_options(row["Allowed Options"]),
            "officialFieldName": str(row["Official Field Name"]),
            "presence": "MANDATORY"
            if str(row["Presence"]).upper().startswith("MANDATORY")
            else "OPTIONAL",
            "nvrRefs": parse_nvr_refs(row["NVR refs"]),
            "classification": str(row["BA SSI Classification"]),
            "ruleId": str(row["Rule ID"]),
            "evidence": str(row["MRG evidence"]),
        }
    if len(result) != 164:
        raise AssertionError(f"expected 164 governed rules, got {len(result)}")
    return result


def update_catalogue(workbook: Any, source: Path, target: Path) -> dict[str, int]:
    catalogue = json.loads(source.read_text(encoding="utf-8"))
    rules = governed_rules(workbook)
    touched = 0
    matched: set[tuple[str, str, str]] = set()
    for mapping in catalogue["mappings"]:
        key = (
            str(mapping.get("messageType", "")),
            str(mapping.get("sequence", "")),
            str(mapping.get("tag", "")),
        )
        rule = rules.get(key)
        if not rule or str(mapping.get("option", "")) not in rule["options"]:
            continue
        mapping["officialFieldName"] = rule["officialFieldName"]
        mapping["presence"] = rule["presence"]
        mapping["nvrRefs"] = rule["nvrRefs"]
        mapping["ruleId"] = rule["ruleId"]
        mapping["governanceSource"] = "MT347_SR2026_SSI_TDD_CONTROLLED_v5.xlsx"
        mapping["governanceSourceSha256"] = WORKBOOK_SHA256
        matched.add(key)
        touched += 1
    missing = sorted(set(rules) - matched)
    if missing:
        raise AssertionError(f"catalogue lacks governed rules: {missing[:10]}")
    catalogue["catalogueVersion"] = "SR2026-MT347-TDD-v5"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(catalogue, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"rules": len(rules), "mappingRecords": touched}


def split_reference(value: Any) -> tuple[str | None, int | None]:
    text = str(value or "").strip()
    identity = re.search(
        r"([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})",
        text,
        re.I,
    )
    version = re.search(r"(?:@|/\s*)v(\d+)\b", text, re.I)
    if not identity:
        return None, None
    return identity.group(1), int(version.group(1)) if version else None


def parse_expected_tags(value: Any, default_sequence: str) -> dict[str, str]:
    return {
        f"{match.group(1) or default_sequence}.{match.group(2).upper()}": match.group(3).strip()
        for match in re.finditer(
            r"\b(?:([A-Z0-9]+)\.)?(5[3-8][A-Z])=([^;]+)",
            str(value or ""),
        )
    }


def parse_expected_present_tags(value: Any, default_sequence: str) -> list[str]:
    text = str(value or "")
    tags = {
        match.group(1).upper()
        for match in re.finditer(r"\b(5[3-8][A-Z])\s+(?:emitted|resolved)\b", text, re.I)
    }
    for match in re.finditer(
        r"\b(5[3-8][A-Z](?:/5[3-8][A-Z])+)\s+(?:emitted|resolved)\b",
        text,
        re.I,
    ):
        tags.update(part.upper() for part in match.group(1).split("/"))
    return [f"{default_sequence}.{tag}" for tag in sorted(tags)]


def profile_context(
    catalogue: dict[str, Any],
    message_type: str,
    business_scenario: str,
    input_contract: str,
    expected_output: str,
) -> tuple[str, str, str]:
    prefix = business_scenario.split(" / ", 1)[0]
    sequence = prefix.split(" - ", 1)[0].strip()
    if sequence.lower() == "message":
        sequence = "MESSAGE"
    message_rows = [
        row for row in catalogue["mappings"] if row.get("messageType") == message_type
    ]
    available_sequences = {str(row.get("sequence")) for row in message_rows}
    if sequence not in available_sequences:
        qualified_sequences = set(
            re.findall(
                r"\b([A-Z][A-Z0-9]*)\.5[3-8][A-Z]\b",
                f"{input_contract};{expected_output}",
            )
        )
        governed_matches = qualified_sequences & available_sequences
        if len(governed_matches) == 1:
            sequence = governed_matches.pop()
        elif len(available_sequences) == 1:
            sequence = next(iter(available_sequences))
        elif not message_rows and message_type not in SCOPE:
            return "MESSAGE", "OUT_OF_SCOPE", "MESSAGE"
        else:
            raise AssertionError(
                f"cannot determine governed sequence for {message_type}/{business_scenario}; "
                f"qualified={sorted(qualified_sequences)}, available={sorted(available_sequences)}"
            )
    candidates = [
        row for row in message_rows if row.get("sequence") == sequence
    ]
    if not candidates:
        raise AssertionError(f"no catalogue profile for {message_type}/{sequence}")
    functions = {str(row["businessFunction"]) for row in candidates}
    legs = {str(row.get("settlementLeg") or sequence) for row in candidates}
    if len(functions) != 1 or len(legs) != 1:
        raise AssertionError(f"ambiguous catalogue profile for {message_type}/{sequence}")
    return sequence, functions.pop(), legs.pop()


def canonical_role_for_tag(
    catalogue: dict[str, Any],
    message_type: str,
    sequence: str,
    tag: str,
    *,
    ssi_supported_only: bool = True,
) -> str | None:
    roles = {
        str(row["canonicalRole"])
        for row in catalogue["mappings"]
        if row.get("messageType") == message_type
        and row.get("sequence") == sequence
        and row.get("tag") == tag
        and (not ssi_supported_only or row.get("scopeStatus") == "SSI_SUPPORTED")
    }
    if len(roles) > 1:
        raise AssertionError(f"ambiguous canonical role for {message_type}/{sequence}/{tag}")
    return next(iter(roles), None)


def add_named_input_roles(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
    ssi_values: dict[str, str],
    transaction_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    beneficiary_is_transaction = "beneficiaryProvenance=TRANSACTION_CONTEXT" in input_contract
    for input_name, tag in input_configuration["namedInputTags"].items():
        match = re.search(rf"\b{re.escape(input_name)}=(BANK-SVC-[A-Z0-9]+)", input_contract)
        if not match:
            continue
        service_id = match.group(1)
        if service_id not in input_configuration["bankServices"]:
            continue
        canonical_role = canonical_role_for_tag(catalogue, message_type, sequence, tag)
        if not canonical_role:
            continue
        transaction_owned = tag == "58" and beneficiary_is_transaction
        target = transaction_values if transaction_owned else ssi_values
        target[canonical_role] = input_configuration["bankServices"][service_id]
        provenance[canonical_role] = {
            "owner": "TRANSACTION_CONTEXT" if transaction_owned else "SSI",
            "sourceType": "BANK_SERVICE_ID",
            "sourceId": service_id,
        }


def add_route_profile_roles(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
    ssi_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    route_match = re.search(r"\broute=(DIRECT|INTERMEDIARY)\b", input_contract)
    if not route_match:
        return
    route_name = route_match.group(1)
    profile = input_configuration["routeProfiles"][route_name]
    for tag, service_id in profile.items():
        canonical_role = canonical_role_for_tag(catalogue, message_type, sequence, tag)
        if canonical_role:
            ssi_values.setdefault(canonical_role, input_configuration["bankServices"][service_id])
            provenance.setdefault(canonical_role, {
                "owner": "SSI",
                "sourceType": "CONTROLLED_ROUTE_PROFILE",
                "sourceId": service_id,
            })


def add_configured_case_roles(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    case_id: str,
    message_type: str,
    sequence: str,
    ssi_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    overrides = input_configuration.get("caseRoleOverrides", {}).get(case_id, {})
    for tag, service_id in overrides.items():
        canonical_role = canonical_role_for_tag(
            catalogue, message_type, sequence, str(tag)
        )
        if not canonical_role:
            raise AssertionError(f"case override tag is not governed: {case_id}/{tag}")
        ssi_values[canonical_role] = input_configuration["bankServices"][service_id]
        provenance[canonical_role] = {
            "owner": "SSI",
            "sourceType": "CONTROLLED_CASE_INPUT",
            "sourceId": service_id,
        }


def add_clearing_directory_role(
    catalogue: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
    ssi_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    clearing_match = re.search(
        r"Party Identifier\s+(//[A-Z0-9]+)\s+plus BIC\s+([A-Z0-9]+)\s+for\s+(?:[A-Z0-9]+\.)?(5[3-8][A-Z])",
        input_contract,
    )
    if not clearing_match:
        return
    tag_option = clearing_match.group(3)
    tag, option = tag_option[:2], tag_option[2:]
    candidate = next((row for row in catalogue["mappings"] if row.get("messageType") == message_type and row.get("sequence") == sequence and row.get("tag") == tag and row.get("option") == option), None)
    if not candidate:
        return
    canonical_role = str(candidate["canonicalRole"])
    ssi_values[canonical_role] = f"{clearing_match.group(1)}\n{clearing_match.group(2)}"
    provenance[canonical_role] = {
        "owner": "SSI",
        "sourceType": "CLEARING_DIRECTORY",
        "sourceId": clearing_match.group(1),
    }


def add_upstream_transaction_role(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
    transaction_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    if "immutable upstream" not in input_contract or "58A" not in input_contract:
        return
    service_id = input_configuration["upstreamTransactionTags"]["58"]
    canonical_role = canonical_role_for_tag(catalogue, message_type, sequence, "58")
    if canonical_role:
        transaction_values[canonical_role] = input_configuration["bankServices"][service_id]
        provenance[canonical_role] = {
            "owner": "TRANSACTION_CONTEXT",
            "sourceType": "IMMUTABLE_UPSTREAM_INSTRUCTION",
            "sourceId": service_id,
        }


def present_input_tags(input_contract: str) -> set[str]:
    tags = {
        match.group(1).upper()
        for match in re.finditer(
            r"\b(5[3-8])[aA]\s+(?:present|own Nostro available)\b",
            input_contract,
        )
    }
    for match in re.finditer(
        r"\bvalid(?:\s+MT\d+)?\s+(5[3-8]A(?:/5[3-8]A)*)",
        input_contract,
        re.I,
    ):
        tags.update(part[:2] for part in match.group(1).upper().split("/"))
    return tags


def add_present_input_roles(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
    ssi_values: dict[str, str],
    transaction_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    defaults = input_configuration["routeProfiles"]["DIRECT"]
    for tag in present_input_tags(input_contract):
        service_id = defaults.get(tag)
        canonical_role = canonical_role_for_tag(
            catalogue, message_type, sequence, tag, ssi_supported_only=False
        )
        if not service_id or not canonical_role:
            continue
        transaction_owned = not any(
            row.get("messageType") == message_type
            and row.get("sequence") == sequence
            and row.get("tag") == tag
            and row.get("scopeStatus") == "SSI_SUPPORTED"
            for row in catalogue["mappings"]
        )
        target = transaction_values if transaction_owned else ssi_values
        target.setdefault(
            canonical_role, input_configuration["bankServices"][service_id]
        )
        provenance.setdefault(canonical_role, {
            "owner": "TRANSACTION_CONTEXT" if transaction_owned else "SSI",
            "sourceType": (
                "IMMUTABLE_UPSTREAM_INSTRUCTION"
                if transaction_owned
                else "CONTROLLED_PRESENT_INPUT"
            ),
            "sourceId": service_id,
        })


def role_values_from_input(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    input_contract: str,
) -> tuple[dict[str, str], dict[str, str], dict[str, dict[str, str]]]:
    ssi_values: dict[str, str] = {}
    transaction_values: dict[str, str] = {}
    provenance: dict[str, dict[str, str]] = {}
    add_named_input_roles(catalogue, input_configuration, message_type, sequence, input_contract, ssi_values, transaction_values, provenance)
    add_route_profile_roles(catalogue, input_configuration, message_type, sequence, input_contract, ssi_values, provenance)
    add_clearing_directory_role(catalogue, message_type, sequence, input_contract, ssi_values, provenance)
    add_upstream_transaction_role(catalogue, input_configuration, message_type, sequence, input_contract, transaction_values, provenance)
    add_present_input_roles(
        catalogue,
        input_configuration,
        message_type,
        sequence,
        input_contract,
        ssi_values,
        transaction_values,
        provenance,
    )
    return ssi_values, transaction_values, provenance


def complete_positive_mandatory_roles(
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
    message_type: str,
    sequence: str,
    ssi_values: dict[str, str],
    provenance: dict[str, dict[str, str]],
) -> None:
    """Complete mandatory route roles from input configuration, never expected output."""
    bank_services = input_configuration["bankServices"]
    defaults = input_configuration["routeProfiles"]["DIRECT"]
    mandatory = {
        str(row["tag"]): str(row["canonicalRole"])
        for row in catalogue["mappings"]
        if row.get("messageType") == message_type
        and row.get("sequence") == sequence
        and row.get("option") == "A"
        and row.get("scopeStatus") == "SSI_SUPPORTED"
        and row.get("presence") == "MANDATORY"
    }
    for tag, role in mandatory.items():
        service_id = defaults.get(tag)
        if role in ssi_values or service_id not in bank_services:
            continue
        ssi_values[role] = bank_services[service_id]
        provenance[role] = {
            "owner": "SSI",
            "sourceType": "CONTROLLED_MANDATORY_ROUTE_DEFAULT",
            "sourceId": service_id,
        }


def fixture_record(
    case: dict[str, Any],
    binding: dict[str, Any],
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
) -> dict[str, Any]:
    ssi_id, ssi_version = split_reference(binding["SSI ID / code / version"])
    applicability_id, applicability_version = split_reference(binding["Applicability ID / version"])
    nostro_id, nostro_version = split_reference(binding["Nostro ID / version"])
    rma_id, rma_version = split_reference(binding["RMA ID / version"])
    fixture_set = str(case["Fixture set"])
    message_type = str(case["Message Type"])
    business_scenario = str(case["Business Scenario"])
    input_contract = str(case["Input"])
    expected_output = str(case["MT Expected Output"])
    sequence, business_function, settlement_leg = profile_context(
        catalogue,
        message_type,
        business_scenario,
        input_contract,
        expected_output,
    )
    expected_tags = parse_expected_tags(expected_output, sequence)
    expected_present_tags = parse_expected_present_tags(expected_output, sequence)
    ssi_role_values, transaction_role_values, role_provenance = role_values_from_input(
        catalogue,
        input_configuration,
        message_type,
        sequence,
        input_contract,
    )
    add_configured_case_roles(
        catalogue,
        input_configuration,
        str(case[CASE_ID_COLUMN]),
        message_type,
        sequence,
        ssi_role_values,
        role_provenance,
    )
    if str(case["Polarity"]) == "POSITIVE":
        complete_positive_mandatory_roles(
            catalogue,
            input_configuration,
            message_type,
            sequence,
            ssi_role_values,
            role_provenance,
        )
    supplemental_account = input_configuration.get("supplementalAccounts", {}).get(
        str(case[CASE_ID_COLUMN])
    )
    hybrid = input_configuration.get("hybridTransactions", {}).get(
        str(case[CASE_ID_COLUMN])
    )
    if hybrid:
        tag = str(hybrid["tag"])
        role = canonical_role_for_tag(
            catalogue,
            message_type,
            sequence,
            tag,
            ssi_supported_only=False,
        )
        service_id = str(hybrid["institutionBankServiceId"])
        institution = input_configuration["bankServices"][service_id]
        account_reference = hybrid.get("accountReference")
        if not role:
            raise AssertionError(
                f"hybrid role is not governed for {case[CASE_ID_COLUMN]}/{tag}"
            )
        rendered = (
            f"/{account_reference}\n{institution}"
            if account_reference
            else institution
        )
        transaction_role_values[role] = rendered
        role_provenance[role] = {
            "owner": "TRANSACTION_CONTEXT",
            "sourceType": "IMMUTABLE_UPSTREAM_INSTRUCTION",
            "sourceId": service_id,
        }
        expected_tags[f"{sequence}.{tag}A"] = rendered
    if supplemental_account and not nostro_id:
        nostro_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"MT347:{case[CASE_ID_COLUMN]}:NOSTRO"))
        nostro_version = 1
    return {
        "testCaseId": str(case[CASE_ID_COLUMN]),
        "bindingId": str(binding["Binding ID / version"]),
        "fixtureSet": fixture_set,
        "messageType": message_type,
        "businessScenario": business_scenario,
        "businessFunction": business_function,
        "sequence": sequence,
        "settlementLeg": settlement_leg,
        "inputContract": input_contract,
        "ownershipContract": str(case["SSI/Nostro ownership and source"]),
        "expectedTags": expected_tags,
        "expectedPresentTags": expected_present_tags,
        "ssiRoleValues": ssi_role_values,
        "transactionRoleValues": transaction_role_values,
        "roleProvenance": role_provenance,
        "roleValuesSource": "INPUT_CONFIGURATION",
        "expectedOutput": expected_output,
        "expectedStatus": str(case["Expected status/error"]),
        "expectedHttp": str(case["Expected HTTP"]),
        "polarity": str(case["Polarity"]),
        "mrgEvidence": str(case["MRG source/page/evidence"]),
        "identity": {
            "ssi": {"id": ssi_id, "version": ssi_version},
            "applicability": {"id": applicability_id, "version": applicability_version},
            "nostro": {"id": nostro_id, "version": nostro_version},
            "rma": {"id": rma_id, "version": rma_version},
        },
        "expectedAccountReference": supplemental_account
        or (
            None
            if str(binding["Expected rendered accountReference"]).startswith("N/A")
            else str(binding["Expected rendered accountReference"])
        ),
        "bankServiceBindings": str(binding["Bank Service IDs"]),
    }


def write_fixture_artifacts(
    workbook: Any,
    output_dir: Path,
    catalogue: dict[str, Any],
    input_configuration: dict[str, Any],
) -> dict[str, Any]:
    cases = {row[CASE_ID_COLUMN]: row for row in rows_by_header(workbook, "TDD Cases")}
    bindings = {row[CASE_ID_COLUMN]: row for row in rows_by_header(workbook, "Fixture Manifest")}
    if set(cases) != set(bindings) or len(cases) != 362:
        raise AssertionError("TDD Cases and Fixture Manifest must have the same 362 identities")
    records = [
        fixture_record(cases[key], bindings[key], catalogue, input_configuration)
        for key in sorted(cases)
    ]
    groups = {
        "positive": [row for row in records if row["polarity"] == "POSITIVE"],
        "negative": [row for row in records if row["polarity"] == "NEGATIVE"],
        "boundary": [row for row in records if row["polarity"] == "BOUNDARY"],
    }
    expected = {"positive": 152, "negative": 208, "boundary": 2}
    actual = {key: len(value) for key, value in groups.items()}
    if actual != expected:
        raise AssertionError(f"unexpected polarity counts: {actual}")
    output_dir.mkdir(parents=True, exist_ok=True)
    artifacts: dict[str, Any] = {}
    for name, rows in groups.items():
        path = output_dir / f"mt347-{name}.v1.json"
        body = {
            "schemaVersion": "1.0",
            "fixtureId": f"MT347-{name.upper()}-V1",
            "classification": "SYNTHETIC_DEMO_QA_UAT",
            "loadPolicy": "CANONICAL_RELOAD" if name == "positive" else "ISOLATED_TEST_ONLY",
            "sourceWorkbook": "qa/tdd/mt347/MT347_SR2026_SSI_TDD_CONTROLLED_v5.xlsx",
            "sourceWorkbookSha256": WORKBOOK_SHA256,
            "memorySha256": MEMORY_SHA256,
            "records": rows,
        }
        path.write_text(json.dumps(body, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        artifacts[name] = {"path": path.as_posix(), "count": len(rows), "sha256": sha256(path)}
    manifest_path = output_dir / "mt347-fixtures.v1.manifest.json"
    manifest = {
        "schemaVersion": "1.0",
        "fixtureFamily": "MT347-SR2026-SSI",
        "sourceWorkbookSha256": WORKBOOK_SHA256,
        "memorySha256": MEMORY_SHA256,
        "artifacts": artifacts,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"counts": actual, "manifest": manifest_path.as_posix(), "manifestSha256": sha256(manifest_path)}


def payload_row(columns: list[str], identity: str, payload: dict[str, Any], updated_at: str) -> list[Any]:
    values = {"id": identity, "payload": canonical_json(payload), "updated_at": updated_at}
    return [values.get(column) for column in columns]


def build_positive_payloads(
    record: dict[str, Any],
    input_configuration: dict[str, Any],
    timestamp: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any], dict[str, Any] | None]:
    identity = record["identity"]
    ssi_id = identity["ssi"]["id"]
    applicability_id = identity["applicability"]["id"]
    rma_id = identity["rma"]["id"]
    if not ssi_id or not applicability_id or not rma_id:
        raise AssertionError(f"positive identity incomplete: {record['testCaseId']}")
    receiver_service_id = input_configuration["messageReceiverBankServices"].get(
        record["messageType"]
    )
    if not receiver_service_id:
        raise AssertionError(f"message Receiver not configured: {record['messageType']}")
    counterparty_bic = input_configuration["bankServices"][receiver_service_id]
    ssi_payload = {
        "id": ssi_id,
        "fixtureFamily": "MT347-SR2026-SSI",
        "fixtureBindingId": record["bindingId"],
        "counterpartyId": f"MT347-{counterparty_bic}",
        "scope": "STANDING",
        "maker": MT347_MAKER,
        "route": {
            "ssiCode": f"SSI-{record['testCaseId']}",
            "messageTypes": [record["messageType"]],
            "businessFunction": record["businessFunction"],
            "sequence": record["sequence"],
            "settlementLeg": record["settlementLeg"],
            "currency": "USD",
            "bookingEntity": "HK01",
            "counterpartyBic": counterparty_bic,
            "roleValues": record["ssiRoleValues"],
            "roleProvenance": {
                role: evidence
                for role, evidence in record["roleProvenance"].items()
                if evidence["owner"] == "SSI"
            },
            "fixtureSet": record["fixtureSet"],
        },
        "status": "ACTIVE",
        "version": identity["ssi"]["version"],
        "createdAt": timestamp,
        "updatedAt": timestamp,
        "checker": MT347_CHECKER,
    }
    applicability_payload = {
        "id": applicability_id,
        "ssiId": ssi_id,
        "fixtureFamily": "MT347-SR2026-SSI",
        "fixtureBindingId": record["bindingId"],
        "consumer": "TREASURY" if record["messageType"].startswith("MT3") else "TRADE_FINANCE",
        "product": record["businessFunction"],
        "businessFunction": record["businessFunction"],
        "paymentLeg": record["settlementLeg"],
        "messageType": record["messageType"],
        "sequence": record["sequence"],
        "settlementLeg": record["settlementLeg"],
        "currency": "USD",
        "direction": "OUTBOUND",
        "status": "ACTIVE",
        "validFrom": "2026-01-01",
        "validTo": "2027-12-31",
        "version": identity["applicability"]["version"],
        "createdAt": timestamp,
        "updatedAt": timestamp,
        "transactionRoleValues": record["transactionRoleValues"],
        "transactionRoleProvenance": {
            role: evidence
            for role, evidence in record["roleProvenance"].items()
            if evidence["owner"] == "TRANSACTION_CONTEXT"
        },
    }
    rma_payload = {
        "id": rma_id,
        "fixtureFamily": "MT347-SR2026-SSI",
        "fixtureBindingId": record["bindingId"],
        "ownBic": "DEMOHKHH",
        "counterpartyBic": counterparty_bic,
        "service": "FIN",
        "direction": "OUTBOUND",
        "messageTypes": [record["messageType"]],
        "validFrom": "2026-01-01",
        "validTo": "2027-12-31",
        "maker": MT347_MAKER,
        "source": "SYNTHETIC_DEMO",
        "status": "ACTIVE",
        "version": identity["rma"]["version"],
        "createdAt": timestamp,
        "updatedAt": timestamp,
        "checker": MT347_CHECKER,
    }
    nostro_id = identity["nostro"]["id"]
    if not nostro_id:
        return ssi_payload, applicability_payload, rma_payload, None
    account_reference = record["expectedAccountReference"]
    nostro_payload = {
        "id": nostro_id,
        "fixtureFamily": "MT347-SR2026-SSI",
        "fixtureBindingId": record["bindingId"],
        "ownLegalEntityId": "HK01",
        "allowedBookingEntities": ["HK01"],
        "accountServicerBic": input_configuration["bankServices"][
            input_configuration["accountServicerBankServiceId"]
        ],
        "currency": "USD",
        "maskedAccountRef": account_reference,
        "accountReference": account_reference,
        "purpose": "SETTLEMENT",
        "priority": 10,
        "validFrom": "2026-01-01",
        "validTo": "2027-12-31",
        "maker": MT347_MAKER,
        "source": "SYNTHETIC_DEMO",
        "status": "ACTIVE",
        "version": identity["nostro"]["version"],
        "createdAt": timestamp,
        "updatedAt": timestamp,
        "checker": MT347_CHECKER,
    }
    return ssi_payload, applicability_payload, rma_payload, nostro_payload


def insert_positive_database(
    root: Path,
    records: list[dict[str, Any]],
    input_configuration: dict[str, Any],
) -> dict[str, Any]:
    output = root / "data/qa/mt347/ssi-demo.mt347-positive.v1.sqlite"
    if output.exists():
        output.unlink()
    for suffix in ("-wal", "-shm"):
        sidecar = Path(f"{output}{suffix}")
        if sidecar.exists():
            sidecar.unlink()
    source = sqlite3.connect(root / "data/ssi-demo.sqlite")
    target = sqlite3.connect(output)
    try:
        source.backup(target)
        timestamp = "2026-09-12T00:00:00.000Z"
        for table in ("ssi_applicability", "ssi", "nostro_account", "rma_authorisation"):
            target.execute(
                f'DELETE FROM "{table}" WHERE payload LIKE ?',
                ('%"fixtureFamily":"MT347-SR2026-SSI"%',),
            )
        for record in records:
            ssi_payload, applicability_payload, rma_payload, nostro_payload = (
                build_positive_payloads(record, input_configuration, timestamp)
            )
            ssi_id = ssi_payload["id"]
            applicability_id = applicability_payload["id"]
            rma_id = rma_payload["id"]
            target.execute(
                "INSERT OR REPLACE INTO ssi(id,payload,updated_at) VALUES(?,?,?)",
                (ssi_id, canonical_json(ssi_payload), timestamp),
            )
            target.execute(
                "INSERT OR REPLACE INTO ssi_applicability(id,ssi_id,payload,updated_at) VALUES(?,?,?,?)",
                (applicability_id, ssi_id, canonical_json(applicability_payload), timestamp),
            )
            target.execute(
                "INSERT OR REPLACE INTO rma_authorisation(id,payload,updated_at) VALUES(?,?,?)",
                (rma_id, canonical_json(rma_payload), timestamp),
            )
            if nostro_payload:
                target.execute(
                    "INSERT OR REPLACE INTO nostro_account(id,payload,updated_at) VALUES(?,?,?)",
                    (nostro_payload["id"], canonical_json(nostro_payload), timestamp),
                )
        target.commit()
        if target.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise AssertionError("MT347 positive SQLite integrity check failed")
    finally:
        source.close()
        target.close()

    rebuild_path = root / "scripts/rebuild-demo-database.py"
    module_spec = importlib.util.spec_from_file_location("rebuild_demo_database", rebuild_path)
    if module_spec is None or module_spec.loader is None:
        raise AssertionError(f"cannot load canonical seed exporter: {rebuild_path}")
    rebuild_module = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(rebuild_module)
    export_result = rebuild_module.export_seed(
        output,
        root
        / "data/qa/mt347/ssi-demo.mt347-positive.v1.canonical.seed.json",
    )
    return {"sqlite": output.as_posix(), "sqliteSha256": sha256(output), "seed": export_result}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    args = parser.parse_args()
    root = args.root.resolve()
    workbook_path = root / "qa/tdd/mt347/MT347_SR2026_SSI_TDD_CONTROLLED_v5.xlsx"
    if sha256(workbook_path) != WORKBOOK_SHA256:
        raise AssertionError("controlled TDD v5 SHA-256 mismatch")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    catalogue_path = root / CATALOGUE_FILE
    catalogue_result = update_catalogue(workbook, catalogue_path, catalogue_path)
    catalogue = json.loads(catalogue_path.read_text(encoding="utf-8"))
    input_configuration = json.loads(
        (root / "parameters/mt347-fixture-inputs.v1.json").read_text(encoding="utf-8")
    )
    if input_configuration.get("sourcePolicy") != "INPUT_ONLY_NO_EXPECTED_ORACLE":
        raise AssertionError("MT347 fixture input configuration must prohibit expected-oracle input")
    fixtures = write_fixture_artifacts(
        workbook,
        root / "data/qa/mt347",
        catalogue,
        input_configuration,
    )
    positive = json.loads(
        (root / "data/qa/mt347/mt347-positive.v1.json").read_text(encoding="utf-8")
    )["records"]
    database = insert_positive_database(root, positive, input_configuration)
    print(json.dumps({"catalogue": catalogue_result, "fixtures": fixtures, "database": database}, indent=2))


if __name__ == "__main__":
    main()
