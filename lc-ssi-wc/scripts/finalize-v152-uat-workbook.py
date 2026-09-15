"""Create the v15.2 R8 UAT close-out workbook without overwriting R7."""
from __future__ import annotations

import copy
import hashlib
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "qa/mt2/uat/MT2xx_SSI_Resolution_UAT執行清單_v15.2_DRAFT_R7.xlsx"
OUTPUT = ROOT / "qa/mt2/uat/MT2xx_SSI_Resolution_UAT執行清單_v15.2_DRAFT_R8.xlsx"
REPORT = ROOT / "qa/mt2/reports/MT2XX_v15.2_api_evidence_20260911.json"
BROWSER_REPORT = ROOT / "qa/mt2/reports/MT2XX_v15.2_browser_UAT_full_postfix_20260911.json"
ADDENDUM = ROOT / "qa/mt2/reports/MT2XX_v15.1_48_CASE_SCOPE_ADDENDUM_ZH_20260911.md"
OVERLAY_REGRESSION = ROOT / "qa/mt2/reports/MT2XX_v15.2_overlay_regression_20260911.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def chosen_text(response: dict) -> str:
    chosen = response.get("chosenRoute") or response.get("mx", {}).get("chosenRoute")
    if not chosen:
        return "null"
    code = chosen.get("ssiCode") or chosen.get("code") or "—"
    version = chosen.get("ssiVersion") or chosen.get("version") or "—"
    return f"{code} / v{version}"


def candidates_text(response: dict) -> str:
    candidates = (
        response.get("candidates")
        or response.get("alternativeRoutes")
        or response.get("mx", {}).get("candidates")
        or []
    )
    if not candidates:
        return "[]"
    return "; ".join(
        f"{item.get('ssiCode', '—')}/v{item.get('ssiVersion', item.get('version', '—'))}"
        for item in candidates
    )


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for column in range(1, ws.max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def main() -> None:
    report = json.loads(REPORT.read_text(encoding="utf-8"))
    assert report["planned"] == report["passed"] == 20 and report["failed"] == 0
    explicit = report["controls"][0]
    assert explicit["controlId"] == "OPEN-006-EXPLICIT-57A" and explicit["status"] == "PASS"
    results = {item["caseId"]: item for item in report["results"]}
    assert len(results) == 20

    workbook = load_workbook(SOURCE)
    sheet = workbook["v15.2 新增案例"]
    headers = {cell.value: cell.column for cell in sheet[4]}
    for row in range(5, 25):
        case_id = sheet.cell(row, headers["案例編號"]).value
        result = results[case_id]
        evidence = json.loads((ROOT / result["evidencePath"]).read_text(encoding="utf-8"))
        response = evidence["response"]
        values = {
            "狀態": "CLOSED",
            "未決項": "—（連結 OPEN 已關閉）",
            "Evidence 狀態": "DELIVERED",
            "Actual HTTP": result["httpStatus"],
            "Actual code": result["code"],
            "Actual reasonCode": result["reasonCode"] or "—",
            "Actual chosen SSI / ver": chosen_text(response),
            "Actual candidates": candidates_text(response),
            "Snapshot Hash": result["snapshotSha256"],
            "Resolution Token": result["resolutionToken"] or "N/A（422 fail-closed）",
            "Correlation / Request ID": result["correlationId"],
            "Tester": evidence["execution"]["tester"],
            "Execution Date": date.fromisoformat(evidence["execution"]["executionDate"]),
            "PASS / FAIL": result["status"],
            "Screenshot / Defect ID": (
                "Browser UAT full report" if case_id == "BOOK-05" else
                "DEF-001 CLOSED" if case_id == "GEN-202-02" else
                "API evidence report"
            ),
            "Evidence SHA-256": result["evidenceSha256"],
        }
        for header, value in values.items():
            sheet.cell(row, headers[header]).value = value
        sheet.cell(row, headers["Execution Date"]).number_format = "yyyy-mm-dd"

    open_sheet = workbook["未決項清單"]
    open_headers = {cell.value: cell.column for cell in open_sheet[4]}
    closures = {
        "OPEN-007": "CLOSED：UI/API 已綁定 scenarioCode、pinned debit/credit nostroId+version 與 receiverBankServiceId；BOOK-01～05 browser/API evidence PASS。",
        "OPEN-008": "CLOSED：evidence schema v1.0、20-case assertion registry、atomic raw-byte SHA runner 已實作；20/20 evidence 驗證通過。",
        "OPEN-003": "CLOSED：GEN-202-02 與 PROV evidence 證明 instructedAgent=CITIUS33，來源 actualReceiverBic；Cdtr=CHASUS33 且來源 REQUEST_PASS_THROUGH。",
        "OPEN-004": f"CLOSED：舊 48/48 報告已由 scope addendum 限縮，補登 baseline snapshot 52A58C…CEEF6；addendum SHA-256={sha256(ADDENDUM)}。",
        "OPEN-006": f"CLOSED：QA-only overlay 明示 57A 控制測試 PASS（57A=BARCGB22；58A=QA-OWN-57A-USD-PRIMARY + DEMOHKHH）；原 48 案逐案 drift=0；regression SHA-256={sha256(OVERLAY_REGRESSION)}。",
    }
    for row in range(5, open_sheet.max_row + 1):
        open_id = open_sheet.cell(row, open_headers["未決項"]).value
        if open_id in closures:
            open_sheet.cell(row, open_headers["狀態"]).value = "CLOSED"
            open_sheet.cell(row, open_headers["內容／證據"]).value = closures[open_id]

    evidence_sheet = workbook["證據與範圍"]
    start = evidence_sheet.max_row + 1
    entries = [
        ("V15.2-R8-UAT", "20 新增案例執行", "20/20 PASS；evidence 完整性 20/20；0 INCOMPLETE", REPORT.relative_to(ROOT).as_posix(), sha256(REPORT)),
        ("V15.2-BROWSER", "瀏覽器 UAT", "20/20 PASS；含 BOOK stale/race 與 counterparty SSI SKIPPED", BROWSER_REPORT.relative_to(ROOT).as_posix(), sha256(BROWSER_REPORT)),
        ("OPEN-006", "明示 57A 控制", "QA-only overlay PASS；57A 明示、58A Sender；baseline 48 案零漂移", OVERLAY_REGRESSION.relative_to(ROOT).as_posix(), sha256(OVERLAY_REGRESSION)),
        ("OPEN-004", "舊報告範圍限縮", "舊 48/48 不再代表 v15.2 agent/BOOK 全域 release gate", ADDENDUM.relative_to(ROOT).as_posix(), sha256(ADDENDUM)),
    ]
    style_row = max(5, evidence_sheet.max_row)
    for offset, entry in enumerate(entries):
        row = start + offset
        copy_row_style(evidence_sheet, style_row, row)
        for column, value in enumerate(entry, 1):
            evidence_sheet.cell(row, column).value = value

    change = workbook.copy_worksheet(workbook["變更說明 R7"])
    change.title = "變更說明 R8"
    for row in change.iter_rows():
        for cell in row:
            cell.value = None
    change["A1"] = "v15.2_DRAFT_R7 → R8 執行與結案說明"
    change["A2"] = "R7 保留不覆寫；R8 回填真實 UAT evidence 並關閉 OPEN-003／004／006／007／008。"
    headers_change = ["#", "項目", "R7", "R8", "證據 / 保護"]
    for column, value in enumerate(headers_change, 1):
        change.cell(4, column).value = value
    changes = [
        (1, "新增 20 案", "BLOCKED / PLANNED", "CLOSED / DELIVERED / PASS 20", f"API report SHA {sha256(REPORT)}"),
        (2, "Browser UAT", "未執行", "20/20 PASS", f"Browser report SHA {sha256(BROWSER_REPORT)}"),
        (3, "OPEN-007／008", "OPEN", "CLOSED", "request binding + evidence infrastructure verified"),
        (4, "OPEN-003", "OPEN", "CLOSED", "GEN-202-02 + PROV raw JSON"),
        (5, "OPEN-004", "OPEN", "CLOSED", "48-case scope addendum + snapshot identity"),
        (6, "OPEN-006", "OPEN", "CLOSED", "explicit 57A control PASS + 48-case zero drift"),
        (7, "原 48 案", "v15.1 controlled matrix", "逐案未改；overlay regression drift=0", "44 RESOLVED + 4 AMBIGUOUS cross-check"),
    ]
    for row_index, row_values in enumerate(changes, 5):
        if row_index > change.max_row:
            copy_row_style(change, 5, row_index)
        for column, value in enumerate(row_values, 1):
            change.cell(row_index, column).value = value

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(OUTPUT)
    print(json.dumps({"path": str(OUTPUT), "sha256": sha256(OUTPUT), "bytes": OUTPUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
