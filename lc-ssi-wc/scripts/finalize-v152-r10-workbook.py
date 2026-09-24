"""Create the v15.2 R10 executed UAT candidate without overwriting controlled R9."""
from __future__ import annotations

import copy
import hashlib
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "qa/reports/latest/mt2/MT2xx_SSI_Resolution_UAT執行清單_v15.2_DRAFT_R9.xlsx"
OUTPUT = ROOT / "qa/reports/latest/mt2/MT2xx_SSI_Resolution_UAT執行清單_v15.2_DRAFT_R10.xlsx"
API_REPORT = ROOT / "qa/reports/latest/mt2/MT2XX_v15.2_api_evidence_20260911.json"
BROWSER_REPORT = ROOT / "qa/reports/latest/mt2/MT2XX_v15.2_browser_UAT_overlay_20260911.json"
OVERLAY_MANIFEST = ROOT / "qa/fixtures/mt2/overlays/ssi-demo.v15.2-own-account-uat.manifest.json"
OVERLAY_REGRESSION = ROOT / "qa/reports/latest/mt2/MT2XX_v15.2_overlay_regression_20260911.json"
SCHEMA = ROOT / "qa/reports/latest/mt2/evidence/v15.2/evidence.schema.json"
REGISTRY = ROOT / "qa/reports/latest/mt2/evidence/v15.2/assertion-registry.json"
VERIFIER = ROOT / "scripts/mt2-v152-evidence-verifier.mjs"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def response_mx(response: dict) -> dict:
    return response.get("mx") or response


def chosen_text(response: dict) -> str:
    chosen = response_mx(response).get("chosenRoute")
    if not chosen:
        return "null"
    code = chosen.get("ssiCode") or chosen.get("code") or "—"
    version = chosen.get("ssiVersion") or chosen.get("version") or "—"
    return f"{code} / v{version}"


def candidates_text(response: dict) -> str:
    mx = response_mx(response)
    candidates = mx.get("candidates") or mx.get("alternativeRoutes") or []
    if not candidates:
        return "[]"
    return "; ".join(
        f"{item.get('ssiCode', '—')}/v{item.get('ssiVersion', item.get('version', '—'))}"
        for item in candidates
    )


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def append_evidence(sheet, rows: list[tuple[str, str, str, str, str]]) -> None:
    start = sheet.max_row + 1
    style_row = max(5, sheet.max_row)
    for offset, values in enumerate(rows):
        row = start + offset
        copy_row_style(sheet, style_row, row)
        for column, value in enumerate(values, 1):
            sheet.cell(row, column).value = value


def main() -> None:
    report = json.loads(API_REPORT.read_text(encoding="utf-8"))
    browser = json.loads(BROWSER_REPORT.read_text(encoding="utf-8"))
    assert report["planned"] == report["passed"] == 20 and report["failed"] == 0
    assert browser["summary"]["planned"] == browser["summary"]["fullyPassed"] == 20
    assert browser["summary"]["failed"] == 0
    assert report["snapshotSha256"] == [
        "371696807835a8c0c1a9ca0bd9bcf4d7d8bd37d457db20404b53c6d772fa8721"
    ]
    explicit = report["controls"][0]
    assert explicit["controlId"] == "OPEN-006-EXPLICIT-57A" and explicit["status"] == "PASS"
    results = {item["caseId"]: item for item in report["results"]}
    assert len(results) == 20

    workbook = load_workbook(SOURCE)
    cases = workbook["v15.2 新增案例"]
    headers = {cell.value: cell.column for cell in cases[4]}
    cases["A1"] = "v15.2_DRAFT_R10 新增案例 — 20 案（執行證據候選版 · NOT APPROVED）"
    cases["A2"] = (
        "20 案已在 PROPOSED_QA_ONLY 隔離 overlay 執行：API 20/20 PASS、Browser 20/20 PASS、"
        "Evidence 20/20 經 schema／registry／raw-byte SHA 與獨立 verifier 核驗。"
    )
    for row in range(5, 25):
        case_id = cases.cell(row, headers["案例編號"]).value
        result = results[case_id]
        evidence_path = ROOT / result["evidencePath"]
        evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
        assert sha256(evidence_path) == result["evidenceSha256"]
        assert evidence["source"]["snapshotIdentityMethod"] == "SQLITE_WAL_AWARE_LOGICAL_SNAPSHOT_V1"
        response = evidence["response"]
        values = {
            "狀態": "CLOSED",
            "未決項": "—（OPEN-003／006／009／010 已關閉）",
            "Evidence 狀態": "DELIVERED",
            "Evidence JSON Path": result["evidencePath"],
            "Actual HTTP": result["httpStatus"],
            "Actual code": result["code"],
            "Actual reasonCode": result["reasonCode"] or "—",
            "Actual chosen SSI / ver": chosen_text(response),
            "Actual candidates": candidates_text(response),
            "Snapshot Hash": result["snapshotSha256"],
            "Resolution Token": result["resolutionToken"] or "N/A（422 fail-closed）",
            "Correlation / Request ID": result["correlationId"],
            "Tester": evidence["execution"]["tester"],
            "Execution Date": date.fromisoformat(evidence["execution"]["executionDate"]),
            "PASS / FAIL": result["status"],
            "Screenshot / Defect ID": (
                f"{rel(BROWSER_REPORT)}; DEF-001 CLOSED"
                if case_id == "GEN-202-02"
                else rel(BROWSER_REPORT)
            ),
            "Evidence SHA-256": result["evidenceSha256"],
        }
        for header, value in values.items():
            cases.cell(row, headers[header]).value = value
        cases.cell(row, headers["Execution Date"]).number_format = "yyyy-mm-dd"

    open_sheet = workbook["未決項清單"]
    open_headers = {cell.value: cell.column for cell in open_sheet[4]}
    open_sheet["A1"] = "未決項清單 — v15.2_DRAFT_R10 單一事實來源"
    open_sheet["A2"] = "OPEN-001～010 皆已關閉；R10 仍為待 BA／QA 簽核的受控 DRAFT，不自動等同 production release。"
    closures = {
        "OPEN-003": "CLOSED：GEN-202-02 與 PROV 證據證明 instructedAgent=CITIUS33（actualReceiverBic）；Cdtr=CHASUS33（REQUEST_PASS_THROUGH）。",
        "OPEN-006": f"CLOSED：QA-only 明示 57A control PASS；overlay manifest SHA-256={sha256(OVERLAY_MANIFEST)}；48-case drift=0。",
        "OPEN-009": "CLOSED：BOOK-01～05 已驗 53B option B／58A 前導斜線／rendered account collision fail-closed／MT-MX 一致；API 與 Browser 皆 PASS。BA 映射：DbtrAcct+SttlmAcct=debit，CdtrAcct=credit。",
        "OPEN-010": f"CLOSED：WAL-aware logical snapshot identity、schema v1.0、全域唯一 registry、raw-byte SHA 與獨立 verifier 全數通過；verifier SHA-256={sha256(VERIFIER)}。",
    }
    for row in range(5, open_sheet.max_row + 1):
        open_id = open_sheet.cell(row, open_headers["未決項"]).value
        if open_id in closures:
            open_sheet.cell(row, open_headers["狀態"]).value = "CLOSED"
            open_sheet.cell(row, open_headers["內容／證據"]).value = closures[open_id]

    evidence_sheet = workbook["證據與範圍"]
    # Preserve the rejected R8 hash as historical evidence; the current API report is a fresh file.
    evidence_sheet["D51"] = "HISTORICAL-R8-HASH-ONLY（不得解析為現行檔案路徑）"
    append_evidence(
        evidence_sheet,
        [
            ("V15.2-R10-API", "20 新增案例 API／Evidence", "20/20 PASS；獨立 verifier 通過；單一 logical snapshot identity。", rel(API_REPORT), sha256(API_REPORT)),
            ("V15.2-R10-BROWSER", "瀏覽器 UAT", "20/20 PASS；BOOK-05 stale/race 來自 raw DOM/network/correlation 重算。", rel(BROWSER_REPORT), sha256(BROWSER_REPORT)),
            ("V15.2-QA-OVERLAY", "PROPOSED_QA_ONLY overlay", "不污染 baseline；manifest 綁定 base snapshot SHA 與 zero-drift regression。", rel(OVERLAY_MANIFEST), sha256(OVERLAY_MANIFEST)),
            ("V15.2-EVIDENCE-SCHEMA", "Evidence schema v1.0", "JSON schema 必填 snapshotIdentityMethod；額外欄位 fail-closed。", rel(SCHEMA), sha256(SCHEMA)),
            ("V15.2-ASSERTION-REGISTRY", "Assertion registry", "caseId 與 assertionId 全域唯一；獨立 verifier 重算。", rel(REGISTRY), sha256(REGISTRY)),
            ("BA-MX-CREDIT-ACCOUNT-MAPPING", "BOOK_TRANSFER MT/MX 帳號映射", "CLOSED — BA CONFIRMED：53B/DbtrAcct=debit；58A/CdtrAcct=credit；INDA SttlmAcct=debit。", "v15.2 新增案例/BOOK-01", sha256(API_REPORT)),
        ],
    )

    change = workbook.copy_worksheet(workbook["變更說明 R9"])
    change.title = "變更說明 R10"
    for row in change.iter_rows():
        for cell in row:
            cell.value = None
    change["A1"] = "v15.2_DRAFT_R9（ACCEPTED CONTROLLED DRAFT）→ R10 執行證據候選"
    change["A2"] = "R9 保留不覆寫；R10 僅回填實際 API／Browser UAT evidence、BA mapping 與 OPEN closure。"
    for column, value in enumerate(["#", "項目", "R9", "R10", "證據 / 保護"], 1):
        change.cell(4, column).value = value
    rows = [
        (1, "新增 20 案", "BLOCKED / PLANNED / 未執行", "CLOSED / DELIVERED / PASS 20", f"API SHA {sha256(API_REPORT)}"),
        (2, "Browser UAT", "待 OPEN-009/010", "20/20 PASS", f"Browser SHA {sha256(BROWSER_REPORT)}"),
        (3, "OPEN-009", "OPEN", "CLOSED", "BOOK message contract API+browser evidence"),
        (4, "OPEN-010", "OPEN", "CLOSED", "independent evidence trust boundary"),
        (5, "OPEN-003/006", "OPEN/controlled", "CLOSED", "DEF-001 provenance + explicit 57A control"),
        (6, "BA MX mapping", "OPEN", "CLOSED — BA CONFIRMED", "DbtrAcct/SttlmAcct=debit; CdtrAcct=credit"),
        (7, "原 48 案", "v15.1 controlled matrix", "A5:AD52 不變；overlay drift=0", f"Regression SHA {sha256(OVERLAY_REGRESSION)}"),
        (8, "假資料隔離", "OPEN-006 fixture pending", "PROPOSED_QA_ONLY overlay", f"Manifest SHA {sha256(OVERLAY_MANIFEST)}"),
    ]
    for row_index, values in enumerate(rows, 5):
        if row_index > change.max_row:
            copy_row_style(change, 5, row_index)
        for column, value in enumerate(values, 1):
            change.cell(row_index, column).value = value
    change.freeze_panes = "A5"

    # The controlled workbook convention is Times New Roman 11 throughout.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(
                    name="Times New Roman",
                    size=11,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    vertAlign=cell.font.vertAlign,
                    underline=cell.font.underline,
                    strike=cell.font.strike,
                    color=copy.copy(cell.font.color),
                )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(OUTPUT)
    print(json.dumps({"path": str(OUTPUT), "sha256": sha256(OUTPUT), "bytes": OUTPUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
