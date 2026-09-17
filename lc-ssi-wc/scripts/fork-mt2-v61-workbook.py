from __future__ import annotations

import hashlib
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "qa/mt2/mt2-final/fixtures/MT2XX_測試案例_SSI與NOSTRO_v6.xlsx"
TARGET = ROOT / "qa/mt2/mt2-final/fixtures/MT2XX_測試案例_SSI與NOSTRO_v6.1_DRAFT.xlsx"
SEED = ROOT / "qa/FIX_DATA/rma/reload-test-data/ssi-demo.v15.8.pacs009-repaired-isolated.canonical.seed.json"

DEBIT = {
    "id": "7520b02e-ec7c-4c50-a9b8-4f78056a05e7",
    "version": 4,
    "accountReference": "DEMO-NOSTRO-001-PRIMARY",
    "accountServicerBic": "CITIUS33",
}
CREDIT = {
    "id": "559335f6-ea8b-49c7-bf86-e7c563f94b19",
    "version": 4,
    "accountReference": "DEMO-NOSTRO-001-EXPCOLL",
    "accountServicerBic": "CITIUS33",
}
RECEIVER_DEBIT = {
    "id": "570a0000-0000-4000-8000-000000000001",
    "version": 1,
    "accountReference": "QA-OWN-57A-USD-PRIMARY",
    "accountServicerBic": "BARCGB22",
}
MT202_14_ROUTE = {
    "ssiId": "a6486f75-a0a8-5207-9c28-cedf30f4f990",
    "ssiCode": "SSI-UAT-GEN-10",
    "ssiVersion": 1,
    "nostroId": "fe5d672f-311e-4eb9-88bf-f8538d45793c",
    "nostroVersion": 4,
    "settlementRouteId": "ROUTE-USD-CITIUS33",
    "accountId": "DEMO-NOSTRO-001-PRIMARY",
    "accountWithBic": "CITIUS33",
}

MT202_RECEIVER_IS_AWI_CASES = {
    "MT202-03",
    "MT202-04",
    "MT202-05",
    "MT202-06",
    "MT202-08",
    "MT202-14",
    "MT202-24",
    "MT202-26",
    "MT202-28",
    "MT202-30",
    "MT202C-01",
    "MT202C-02",
    "MT202C-03",
    "MT202C-04",
    "MT202C-13",
    "MT202C-15",
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def dump(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def migrate_request(raw: str, case_id: str) -> dict:
    value = json.loads(raw)
    request = value["request/context"]
    for key in (
        "ownAccountSubScenario",
        "debitAccount",
        "creditAccount",
        "ownCreditAccount",
        "57A",
        "A.57A",
        "A.53B",
        "A.58A",
        "senderBic",
        "beneficiarySource",
        "skipCounterpartySsiResolution",
    ):
        request.pop(key, None)
    book = case_id in {"MT202-07", "MT202C-17"}
    request.update(
        {
            "scenarioCode": (
                "BOOK_TRANSFER_SAME_RECEIVER"
                if book
                else "CREDIT_ONE_OF_SEVERAL_AT_57A"
            ),
            "ownCreditAccountId": CREDIT["id"],
            "ownCreditAccountVersion": CREDIT["version"],
            "receiverBankServiceId": (
                "BANK-SVC-CITIUS33" if book else "BANK-SVC-BARCGB22"
            ),
        }
    )
    debit = DEBIT if book else RECEIVER_DEBIT
    request.update(
        {
            "ownDebitAccountId": debit["id"],
            "ownDebitAccountVersion": debit["version"],
        }
    )
    value["contractBasis"] = (
        "v15.3 controlled own-account API contract; pinned Nostro id + version"
    )
    value["current DB fixtures"] = [debit, CREDIT]
    return value


def expected(case_id: str) -> tuple[dict, dict]:
    cover = case_id.startswith("MT202C-")
    book = case_id in {"MT202-07", "MT202C-17"}
    scenario = (
        "BOOK_TRANSFER_SAME_RECEIVER"
        if book
        else "CREDIT_ONE_OF_SEVERAL_AT_57A"
    )
    receiver = "CITIUS33" if book else "BARCGB22"
    prefix = "A." if cover else ""
    mx = {
        "httpStatus": 200,
        "decision": "RESOLVED",
        "code": "RESOLVED",
        "resolutionDomain": "OWN_SSI_NOSTRO",
        "counterpartySsiResolution": "SKIPPED",
        "chosenRoute": None,
        "candidates": [],
        "messageDefinitionId": "pacs.009.001.08",
        "businessService": (
            "swift.cbprplus.cov.04" if cover else "swift.cbprplus.04"
        ),
        "canonicalScenario": scenario,
        "canonicalRoles": {
            "sender": "DEMOHKHH",
            "debtor": "DEMOHKHH",
            "receiver": receiver,
            "beneficiary": "DEMOHKHH",
        },
        "resolutionTrace": {"counterpartySsiQueried": False},
        "payloadGenerated": True,
    }
    tags = {f"{prefix}58A": f"/{CREDIT['accountReference']}\nDEMOHKHH"}
    debit = DEBIT if book else RECEIVER_DEBIT
    tags[f"{prefix}53B"] = f"/{debit['accountReference']}"
    omitted: list[str] = []
    if book:
        tags[f"{prefix}53B"] = f"/{DEBIT['accountReference']}"
        omitted.append(f"{prefix}57A")
    else:
        tags[f"{prefix}57A"] = CREDIT["accountServicerBic"]
    mt = {
        "renderer": "MT compatibility view from the same confirmed own-account snapshot",
        "tags": tags,
        "omitted": omitted,
    }
    return mx, mt


def apply_mt202_14_route(raw: str) -> str:
    value = json.loads(raw)
    roles = value.setdefault("canonicalRoles", {})
    roles["selectedSsi"] = MT202_14_ROUTE["ssiCode"]
    value["chosenRoute"] = {
        key: MT202_14_ROUTE[key]
        for key in ("ssiCode", "ssiVersion", "nostroId", "nostroVersion")
    }
    return dump(value)


def apply_mt202_14_fixture(raw: str) -> str:
    value = json.loads(raw)
    fixtures = value.get("current DB fixtures", [])
    replacement = {
        "ssiCode": MT202_14_ROUTE["ssiCode"],
        "ssiId": MT202_14_ROUTE["ssiId"],
        "ssiVersion": MT202_14_ROUTE["ssiVersion"],
        "counterpartyId": "CP-CHASUS33",
        "currency": "USD",
        "accountWithBic": MT202_14_ROUTE["accountWithBic"],
        "accountId": MT202_14_ROUTE["accountId"],
        "nostroId": MT202_14_ROUTE["nostroId"],
        "nostroVersion": MT202_14_ROUTE["nostroVersion"],
        "settlementRouteId": MT202_14_ROUTE["settlementRouteId"],
        "routePurpose": "INTERBANK_TRANSFER",
        "routePreference": "PRIMARY",
        "priority": "10",
    }
    value["current DB fixtures"] = [replacement, *fixtures[1:]]
    return dump(value)


def apply_mt202_receiver_is_awi_rule(raw: str, case_id: str) -> str:
    value = json.loads(raw)
    cover = case_id.startswith("MT202C-")
    tag = "A.57A" if cover else "57A"
    omitted = "A.57a" if cover else "57a"
    value.setdefault("tags", {}).pop(tag, None)
    omitted_fields = value.setdefault("omitted", [])
    if omitted not in omitted_fields:
        omitted_fields.append(omitted)
    return dump(value)


def main() -> None:
    source_sha = digest(SOURCE)
    seed_sha = digest(SEED)
    workbook = load_workbook(SOURCE)
    sheet = workbook["Test Cases"]
    canonical_replacements = 0
    migrated: list[str] = []
    omission_rule_updates: list[str] = []

    for row in range(2, sheet.max_row + 1):
        raw = sheet.cell(row, 4).value
        if isinstance(raw, str) and '"canonicalRoute/agent roles"' in raw:
            sheet.cell(row, 4).value = raw.replace(
                '"canonicalRoute/agent roles"', '"canonicalRoles"'
            )
            canonical_replacements += 1
        case_id = sheet.cell(row, 1).value
        if case_id == "MT202-14":
            sheet.cell(row, 3).value = apply_mt202_14_fixture(
                sheet.cell(row, 3).value
            )
            sheet.cell(row, 4).value = apply_mt202_14_route(
                sheet.cell(row, 4).value
            )
        if case_id in MT202_RECEIVER_IS_AWI_CASES:
            sheet.cell(row, 5).value = apply_mt202_receiver_is_awi_rule(
                sheet.cell(row, 5).value, case_id
            )
            omission_rule_updates.append(case_id)
        if case_id in {"MT202-07", "MT202-32", "MT202C-16", "MT202C-17"}:
            sheet.cell(row, 3).value = dump(
                migrate_request(sheet.cell(row, 3).value, case_id)
            )
            mx, mt = expected(case_id)
            sheet.cell(row, 4).value = dump(mx)
            sheet.cell(row, 5).value = dump(mt)
            migrated.append(case_id)

    if (
        canonical_replacements == 0
        or len(migrated) != 4
        or len(omission_rule_updates) != len(MT202_RECEIVER_IS_AWI_CASES)
    ):
        raise RuntimeError(
            "unexpected migration count: "
            f"canonical={canonical_replacements}, cases={migrated}, "
            f"omissionRule={omission_rule_updates}"
        )

    revision = workbook.create_sheet("Revision v6.1")
    revision.append(["Control", "Value"])
    revision.append(["Status", "DRAFT — NOT APPROVED"])
    revision.append(["Base workbook", SOURCE.name])
    revision.append(["Base SHA-256", source_sha])
    revision.append(["Canonical positive seed", str(SEED.relative_to(ROOT))])
    revision.append(["Canonical seed SHA-256", seed_sha])
    revision.append(["Legacy canonical key replacements", canonical_replacements])
    revision.append(["Own-account contract migrations", ", ".join(migrated)])
    revision.append(
        [
            "MT202 57a receiver-is-AWI corrections",
            ", ".join(omission_rule_updates),
        ]
    )
    revision.append(
        [
            "Negative data rule",
            "Negative overlays remain isolated; never load into the canonical positive database.",
        ]
    )
    revision.freeze_panes = "A2"
    revision.auto_filter.ref = f"A1:B{revision.max_row}"
    revision.column_dimensions["A"].width = 36
    revision.column_dimensions["B"].width = 96

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = "Times New Roman"
                font.sz = 11
                cell.font = font

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(TARGET)
    print(
        json.dumps(
            {
                "target": str(TARGET),
                "baseSha256": source_sha,
                "seedSha256": seed_sha,
                "targetSha256": digest(TARGET),
                "canonicalKeyReplacements": canonical_replacements,
                "migratedCases": migrated,
                "mt202ReceiverIsAwiCases": omission_rule_updates,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
